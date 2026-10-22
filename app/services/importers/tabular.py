"""CSV and Excel (.xlsx) files with a header row, mapped column by column."""
import csv
import io

from openpyxl import load_workbook

from app.services.importers.fields import IMPORT_FIELDS
from app.services.importers.normalize import strip_accents
from app.services.importers.records import MAX_RECORDS, ImportParseError

DELIMITERS = ",;\t|"
PREVIEW_ROWS = 5


def _decode(content: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("latin-1")


def _read_csv(content: bytes) -> list[list[str]]:
    text = _decode(content)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=DELIMITERS)
    except csv.Error:
        dialect = csv.excel
    return [[cell for cell in row] for row in csv.reader(io.StringIO(text), dialect)]


def _read_xlsx(content: bytes) -> list[list[str]]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as err:  # openpyxl raises several unrelated types for bad files
        raise ImportParseError("No se pudo leer el archivo Excel. ¿Es un .xlsx válido?") from err
    sheet = workbook.worksheets[0]
    rows: list[list[str]] = []
    for row in sheet.iter_rows(values_only=True):
        rows.append(["" if cell is None else str(cell) for cell in row])
        if len(rows) > MAX_RECORDS + 1:
            break
    workbook.close()
    return rows


def read_table(content: bytes, filename: str) -> tuple[list[str], list[dict[str, str]]]:
    """Header names and one {header: cell} dict per data row, blank rows dropped."""
    lowered = (filename or "").lower()
    if lowered.endswith(".xls"):
        raise ImportParseError("Los .xls viejos no se pueden leer: guardalo como .xlsx o CSV.")
    is_xlsx = lowered.endswith((".xlsx", ".xlsm")) or content[:2] == b"PK"
    rows = _read_xlsx(content) if is_xlsx else _read_csv(content)

    rows = [row for row in rows if any(cell.strip() for cell in row)]
    if len(rows) < 2:
        raise ImportParseError("El archivo no tiene filas de datos (la primera fila debe ser el encabezado).")

    headers: list[str] = []
    for index, cell in enumerate(rows[0]):
        name = cell.strip() or f"Columna {index + 1}"
        while name in headers:
            name += " (2)"
        headers.append(name)

    data = rows[1:]
    if len(data) > MAX_RECORDS:
        raise ImportParseError(f"El archivo tiene más de {MAX_RECORDS} filas. Dividilo en partes.")
    return headers, [dict(zip(headers, row + [""] * (len(headers) - len(row)))) for row in data]


def _comparable(text: str) -> str:
    return strip_accents(text).lower().replace("_", " ").replace("-", " ").strip()


def suggest_mapping(headers: list[str]) -> dict[str, str | None]:
    """header → field key, by exact alias match first and then by containment.
    Each field is suggested for at most one column."""
    taken: set[str] = set()
    mapping: dict[str, str | None] = {header: None for header in headers}
    normalized = {header: _comparable(header) for header in headers}

    for exact in (True, False):
        for header in headers:
            if mapping[header] is not None:
                continue
            name = normalized[header]
            for field in IMPORT_FIELDS:
                if field.key in taken:
                    continue
                aliases = {_comparable(alias) for alias in (field.key, field.label, *field.aliases)}
                hit = name in aliases if exact else any(len(a) > 3 and a in name for a in aliases)
                if hit:
                    mapping[header] = field.key
                    taken.add(field.key)
                    break
    return mapping


def records_from_table(rows: list[dict[str, str]], mapping: dict[str, str | None]) -> list[dict]:
    """One raw record per row, with only the mapped columns."""
    columns = [(header, key) for header, key in mapping.items() if key]
    return [{key: row.get(header, "") for header, key in columns} for row in rows]
